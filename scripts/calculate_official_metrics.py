"""Recalcula métricas do projeto diretamente nas planilhas oficiais da Receita Federal."""

from __future__ import annotations

import hashlib
import json
import tempfile
import urllib.request
from collections import defaultdict
from decimal import Decimal
from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook

SOURCES = {
    "01b": (
        "https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/publicacoes/"
        "estudos/pessoas-juridicas-por-setor/estudos-setoriais-das-pessoas-juridicas/"
        "dados-setoriais-2024/tabela-01b-secao-sn-e-mei/@@download/file"
    ),
    "05b": (
        "https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/publicacoes/"
        "estudos/pessoas-juridicas-por-setor/estudos-setoriais-das-pessoas-juridicas/"
        "dados-setoriais-2024/tabela-05b-subclasse-sn-e-mei/@@download/file"
    ),
}

MEI = "SIMPLES - MEI"


def download(url: str, destination: Path) -> None:
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(request, timeout=120) as response:
        destination.write_bytes(response.read())


def decimal(value: object) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    return Decimal(str(value))


def integer_or_none(value: object) -> int | None:
    if value in (None, "", "(x)"):
        return None
    return int(value)


def iter_records(path: Path, sheet_name: str) -> Iterator[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value) if value is not None else "" for value in next(rows)]
    for row in rows:
        yield dict(zip(headers, row, strict=False))


def suppressed_quantity(path: Path, sheet_name: str) -> int:
    for record in iter_records(path, sheet_name):
        if record.get("Forma_Tributacao") == MEI:
            return int(record["Quantidade_de_CNPJ"])
    raise RuntimeError(f"Linha {MEI!r} não encontrada em {sheet_name}")


def top(mapping: dict[str, Decimal | int], limit: int = 10) -> list[tuple[str, Decimal | int]]:
    return sorted(mapping.items(), key=lambda item: item[1], reverse=True)[:limit]


def calculate_01b(path: Path) -> dict[str, object]:
    disclosed_quantity = 0
    suppressed_cells = 0
    revenue = Decimal("0")
    mei_collection = Decimal("0")
    collection_all_tax_forms = Decimal("0")
    quantity_by_uf: dict[str, int] = defaultdict(int)
    revenue_by_uf: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))

    for record in iter_records(path, "Secao"):
        collection_all_tax_forms += decimal(record.get("Arrecadacao_MEI_DAS_MEI"))
        if record.get("Forma_Tributacao") != MEI:
            continue

        quantity = integer_or_none(record.get("Quantidade_de_CNPJ"))
        if quantity is None:
            suppressed_cells += 1
        else:
            disclosed_quantity += quantity
            quantity_by_uf[str(record["UF"])] += quantity

        row_revenue = decimal(record.get("Receita_Bruta"))
        revenue += row_revenue
        revenue_by_uf[str(record["UF"])] += row_revenue
        mei_collection += decimal(record.get("Arrecadacao_MEI_DAS_MEI"))

    hidden_quantity = suppressed_quantity(path, "Secao < 4")
    exact_quantity = disclosed_quantity + hidden_quantity

    return {
        "disclosed_quantity": disclosed_quantity,
        "suppressed_quantity": hidden_quantity,
        "exact_quantity": exact_quantity,
        "suppressed_cells": suppressed_cells,
        "revenue": revenue,
        "revenue_per_disclosed_cnpj": revenue / disclosed_quantity,
        "revenue_per_exact_cnpj": revenue / exact_quantity,
        "mei_collection_mei_rows": mei_collection,
        "mei_collection_all_rows": collection_all_tax_forms,
        "top_quantity_uf": top(quantity_by_uf),
        "top_revenue_uf": top(revenue_by_uf),
    }


def calculate_05b(path: Path) -> dict[str, object]:
    disclosed_quantity = 0
    suppressed_cells = 0
    revenue = Decimal("0")
    quantity_by_activity: dict[str, int] = defaultdict(int)
    revenue_by_activity: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))

    for record in iter_records(path, "Subclasse"):
        if record.get("Forma_Tributacao") != MEI:
            continue

        description = str(record.get("Sublasse_CNAE_Descricao") or "")
        quantity = integer_or_none(record.get("Quantidade_de_CNPJ"))
        if quantity is None:
            suppressed_cells += 1
        else:
            disclosed_quantity += quantity
            quantity_by_activity[description] += quantity

        row_revenue = decimal(record.get("Receita_Bruta"))
        revenue += row_revenue
        revenue_by_activity[description] += row_revenue

    hidden_quantity = suppressed_quantity(path, "Subclasse < 4")

    return {
        "disclosed_quantity": disclosed_quantity,
        "suppressed_quantity": hidden_quantity,
        "exact_quantity": disclosed_quantity + hidden_quantity,
        "suppressed_cells": suppressed_cells,
        "revenue": revenue,
        "top_quantity_activity": top(quantity_by_activity),
        "top_revenue_activity": top(revenue_by_activity),
    }


def serialize(value: object) -> object:
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, tuple):
        return [serialize(item) for item in value]
    if isinstance(value, list):
        return [serialize(item) for item in value]
    if isinstance(value, dict):
        return {key: serialize(item) for key, item in value.items()}
    return value


def main() -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        files: dict[str, Path] = {}
        for label, url in SOURCES.items():
            path = root / f"{label}.xlsx"
            download(url, path)
            files[label] = path

        result = {
            "source_sha256": {
                label: hashlib.sha256(path.read_bytes()).hexdigest()
                for label, path in files.items()
            },
            "01b": calculate_01b(files["01b"]),
            "05b": calculate_05b(files["05b"]),
        }
        print(json.dumps(serialize(result), indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
