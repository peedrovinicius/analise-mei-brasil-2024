"""Inspeção controlada das planilhas oficiais Dados Setoriais 2024."""

from __future__ import annotations

import tempfile
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

SOURCES = {
    "01b": (
        "https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/publicacoes/"
        "estudos/pessoas-juridicas-por-setor/estudos-setoriais-das-pessoas-juridicas/"
        "dados-setoriais-2024/tabela-01b-secao-sn-e-mei/@@download/file"
    ),
    "05b": (
        "https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/publicacoes/"
        "estudos/pessoas-juridicas-por-setor/estudos-setoriais-das-pessoas-juridicas/"
        "dados-setoriais-2024/tabela-05b-subclasse-sn-e-mei/@@download/file"
    ),
}


def download(url: str, destination: Path) -> None:
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(request, timeout=60) as response:
        destination.write_bytes(response.read())


def compact(value: object) -> str:
    text = "" if value is None else str(value)
    return text.replace("\n", " ")[:120]


def inspect_workbook(label: str, path: Path) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    print(f"=== {label} ===")
    print(f"arquivo_bytes={path.stat().st_size}")
    print(f"planilhas={workbook.sheetnames}")
    for sheet in workbook.worksheets:
        print(f"--- planilha: {sheet.title} ---")
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = [compact(value) for value in row[:20]]
            print(f"{row_number}: {values}")
            if row_number >= 12:
                break


def main() -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        for label, url in SOURCES.items():
            path = root / f"{label}.xlsx"
            download(url, path)
            inspect_workbook(label, path)


if __name__ == "__main__":
    main()
