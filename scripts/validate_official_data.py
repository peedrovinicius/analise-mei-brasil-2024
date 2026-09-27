"""Valida os resultados publicados contra as planilhas oficiais da Receita Federal."""

from __future__ import annotations

import hashlib
import tempfile
import urllib.request
from collections import defaultdict
from decimal import Decimal
from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook

MEI = "SIMPLES - MEI"

SOURCES = {
    "01b": (
        "https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/publicacoes/"
        "estudos/pessoas-juridicas-por-setor/estudos-setoriais-das-pessoas-juridicas/"
        "dados-setoriais-2024/tabela-01b-secao-sn-e-mei/@@download/file"
    ),
    "05b": (
        "https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/publicacoes/"
        "estudos/pessoas-juridicas-por-setor/estudos-setoriais-das-pessoas-juridicas/"
        "dados-setoriais-2024/tabela-05b-subclasse-sn-e-mei/@@download/file"
    ),
}

EXPECTED_SHA256 = {
    "01b": "493243210f234b8b60548c62b8517d9eac4d0b863d6eeb58a9a59fc656343257",
    "05b": "1bee383882faa8235d070eb41734e125f70469479aa0a79f49b690984b94d5cf",
}

EXPECTED_REVENUE = Decimal("310969612070.00")
EXPECTED_COLLECTION = Decimal("13502226249.57")
EXPECTED_EXACT_QUANTITY = 10_284_140

EXPECTED_01B = {
    "disclosed_quantity": 10_284_095,
    "suppressed_quantity": 45,
    "suppressed_cells": 31,
}

EXPECTED_05B = {
    "disclosed_quantity": 10_277_245,
    "suppressed_quantity": 6_895,
    "suppressed_cells": 4_332,
}

EXPECTED_TOP_QUANTITY_UF = [
    ("SP", 2_877_357),
    ("MG", 1_247_636),
    ("RJ", 940_544),
]

EXPECTED_TOP_REVENUE_UF = [
    ("SP", Decimal("80109942531.61")),
    ("MG", Decimal("41720965574.28")),
    ("PR", Decimal("24118972979.77")),
]

EXPECTED_TOP_QUANTITY_ACTIVITY = [
    ("Cabeleireiros", 640_986),
    ("Comércio varejista de artigos do vestuário e acessórios", 595_465),
    ("Promoção de vendas", 449_990),
    ("Obras de alvenaria", 389_136),
    ("Preparação docum., serv. apoio administrativo não especific.", 372_042),
]

EXPECTED_TOP_REVENUE_ACTIVITY = [
    ("Cabeleireiros", Decimal("19512836744.45")),
    ("Comércio varejista de artigos do vestuário e acessórios", Decimal("17703548358.48")),
    ("Promoção de vendas", Decimal("12490460173.44")),
    ("Obras de alvenaria", Decimal("11583020840.55")),
    ("Preparação docum., serv. apoio administrativo não especific.", Decimal("11326540831.46")),
]


def download(url: str, destination: Path) -> None:
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(request, timeout=120) as response:
        destination.write_bytes(response.read())


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def decimal(value: object) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    return Decimal(str(value))


def quantity_or_none(value: object) -> int | None:
    if value in (None, "", "(x)"):
        return None
    return int(value)


def iter_records(path: Path, sheet_name: str) -> Iterator[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        headers = [str(value) if value is not None else "" for value in next(rows)]
        for row in rows:
            yield dict(zip(headers, row, strict=False))
    finally:
        workbook.close()


def supplemental_quantity(path: Path, sheet_name: str) -> int:
    for record in iter_records(path, sheet_name):
        if record.get("Forma_Tributacao") == MEI:
            return int(record["Quantidade_de_CNPJ"])
    raise AssertionError(f"Linha {MEI!r} não encontrada em {sheet_name!r}.")


def top(mapping: dict[str, Decimal | int], limit: int) -> list[tuple[str, Decimal | int]]:
    return sorted(mapping.items(), key=lambda item: item[1], reverse=True)[:limit]


def calculate_01b(path: Path) -> dict[str, object]:
    disclosed_quantity = 0
    suppressed_cells = 0
    revenue = Decimal("0")
    collection_mei_rows = Decimal("0")
    collection_all_rows = Decimal("0")
    quantity_by_uf: dict[str, int] = defaultdict(int)
    revenue_by_uf: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))

    for record in iter_records(path, "Secao"):
        collection_all_rows += decimal(record.get("Arrecadacao_MEI_DAS_MEI"))
        if record.get("Forma_Tributacao") != MEI:
            continue
        quantity = quantity_or_none(record.get("Quantidade_de_CNPJ"))
        if quantity is None:
            suppressed_cells += 1
        else:
            disclosed_quantity += quantity
            quantity_by_uf[str(record["UF"])] += quantity
        row_revenue = decimal(record.get("Receita_Bruta"))
        revenue += row_revenue
        revenue_by_uf[str(record["UF"])] += row_revenue
        collection_mei_rows += decimal(record.get("Arrecadacao_MEI_DAS_MEI"))

    suppressed_quantity = supplemental_quantity(path, "Secao < 4")
    return {
        "disclosed_quantity": disclosed_quantity,
        "suppressed_quantity": suppressed_quantity,
        "exact_quantity": disclosed_quantity + suppressed_quantity,
        "suppressed_cells": suppressed_cells,
        "revenue": revenue,
        "collection_mei_rows": collection_mei_rows,
        "collection_all_rows": collection_all_rows,
        "top_quantity_uf": top(quantity_by_uf, 3),
        "top_revenue_uf": top(revenue_by_uf, 3),
    }


def calculate_05b(path: Path) -> dict[str, object]:
    disclosed_quantity = 0
    suppressed_cells = 0
    revenue = Decimal("0")
    quantity_by_activity: dict[str, int] = defaultdict(int)
    revenue_by_activity: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))

    for record in iter_records(path, "Subclasse"):
        if record.get("Forma_Tributacao") != MEI:
            continue
        description = str(record.get("Sublasse_CNAE_Descricao") or "")
        quantity = quantity_or_none(record.get("Quantidade_de_CNPJ"))
        if quantity is None:
            suppressed_cells += 1
        else:
            disclosed_quantity += quantity
            quantity_by_activity[description] += quantity
        row_revenue = decimal(record.get("Receita_Bruta"))
        revenue += row_revenue
        revenue_by_activity[description] += row_revenue

    suppressed_quantity = supplemental_quantity(path, "Subclasse < 4")
    return {
        "disclosed_quantity": disclosed_quantity,
        "suppressed_quantity": suppressed_quantity,
        "exact_quantity": disclosed_quantity + suppressed_quantity,
        "suppressed_cells": suppressed_cells,
        "revenue": revenue,
        "top_quantity_activity": top(quantity_by_activity, 5),
        "top_revenue_activity": top(revenue_by_activity, 5),
    }


def assert_equal(name: str, actual: object, expected: object) -> None:
    if actual != expected:
        raise AssertionError(f"{name}: esperado {expected!r}, obtido {actual!r}")


def validate_source_files(files: dict[str, Path]) -> None:
    for label, path in files.items():
        assert_equal(f"sha256_{label}", sha256_file(path), EXPECTED_SHA256[label])

    result_01b = calculate_01b(files["01b"])
    result_05b = calculate_05b(files["05b"])

    for key, expected in EXPECTED_01B.items():
        assert_equal(f"01b_{key}", result_01b[key], expected)
    for key, expected in EXPECTED_05B.items():
        assert_equal(f"05b_{key}", result_05b[key], expected)

    assert_equal("01b_exact_quantity", result_01b["exact_quantity"], EXPECTED_EXACT_QUANTITY)
    assert_equal("05b_exact_quantity", result_05b["exact_quantity"], EXPECTED_EXACT_QUANTITY)
    assert_equal("01b_revenue", result_01b["revenue"], EXPECTED_REVENUE)
    assert_equal("05b_revenue", result_05b["revenue"], EXPECTED_REVENUE)
    assert_equal("collection_mei_rows", result_01b["collection_mei_rows"], EXPECTED_COLLECTION)
    assert_equal("collection_all_rows", result_01b["collection_all_rows"], EXPECTED_COLLECTION)
    assert_equal("top_quantity_uf", result_01b["top_quantity_uf"], EXPECTED_TOP_QUANTITY_UF)
    assert_equal("top_revenue_uf", result_01b["top_revenue_uf"], EXPECTED_TOP_REVENUE_UF)
    assert_equal("top_quantity_activity", result_05b["top_quantity_activity"], EXPECTED_TOP_QUANTITY_ACTIVITY)
    assert_equal("top_revenue_activity", result_05b["top_revenue_activity"], EXPECTED_TOP_REVENUE_ACTIVITY)


def main() -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        files: dict[str, Path] = {}
        for label, url in SOURCES.items():
            path = root / f"{label}.xlsx"
            download(url, path)
            files[label] = path
        validate_source_files(files)

    print("Fonte oficial e resultados publicados validados com sucesso.")
    print(f"CNPJs MEI reconciliados: {EXPECTED_EXACT_QUANTITY:,}")
    print(f"Receita Bruta: R$ {EXPECTED_REVENUE}")
    print(f"Arrecadação DAS-MEI: R$ {EXPECTED_COLLECTION}")


if __name__ == "__main__":
    main()
